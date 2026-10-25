from io import BytesIO
import openpyxl
import pandas as pd
from valuation.data import demo_company, demo_peers, peer_record, BILLION
from valuation.model import (
    BASE_FIELDS, infer_assumptions, default_drivers, forecast, dcf, sensitivity, relative_values
)
from valuation.export import workbook


def test_populated_valuation_results_excel_export():
    c = demo_company()
    historical = c["historical"]
    period = str(historical.columns[-1])
    first_year = pd.Timestamp(period).year + 1
    
    b = {k: float(historical.loc[k, period]) if k in historical.index and pd.notna(historical.loc[k, period]) else 0.0 for k in BASE_FIELDS}
    a = infer_assumptions(historical, b, c)
    drivers = default_drivers(a, first_year)
    f = forecast(b, a, first_year, drivers)
    shares = (c["shares"] or 0) / BILLION
    res = dcf(f, b, a, shares)
    grid = sensitivity(f, b, a, shares)
    
    raw_peers = demo_peers()
    peers = [peer_record(p, period) for p in raw_peers]
    target = dict(
        revenue=b["revenue"],
        ebit=b["revenue"] - b["cogs"] - b["opex"] - b["da"],
        ebitda=b["revenue"] - b["cogs"] - b["opex"],
        net_income=b["net_income"],
        book_equity=max(0.0, b.get("equity", 0.0) - a.other_claims),
    )
    relative = relative_values(peers, target, b["cash"], b["short_debt"] + b["long_debt"], shares, a.other_claims, a.nonoperating_assets)
    
    excel_bytes = workbook(c, historical, b, a, drivers, f, res, grid, peers, relative)
    assert excel_bytes is not None and len(excel_bytes) > 5000
    
    # Load back with openpyxl in data mode
    wb = openpyxl.load_workbook(BytesIO(excel_bytes), data_only=True)
    sheet_names = wb.sheetnames
    
    # 1. Verify all 6 core sheets are present
    expected_sheets = [
        "Dashboard", "Assumptions", "Pro-Forma Forecast", "DCF Bridge", "Relative Valuation", "Historical & Inputs"
    ]
    for s in expected_sheets:
        assert s in sheet_names, f"Missing sheet: {s}"
        
    # 2. Verify Dashboard pre-populated results and recommendation
    ws_dash = wb["Dashboard"]
    blended_price = ws_dash["C27"].value
    ref_price = ws_dash["C28"].value
    upside = ws_dash["C29"].value
    rating = ws_dash["C31"].value
    
    assert isinstance(blended_price, (int, float)) and blended_price > 0, f"Expected numeric blended target price, got: {blended_price}"
    assert isinstance(ref_price, (int, float)) and ref_price > 0, f"Expected numeric reference price, got: {ref_price}"
    assert isinstance(upside, (int, float)), f"Expected numeric upside, got: {upside}"
    assert isinstance(rating, str) and any(r in rating for r in ("BUY", "HOLD", "SELL", "NOT ELIGIBLE")), f"Expected valid recommendation rating, got: {rating}"
    
    # Verify formula & decision rule descriptions in Dashboard
    assert ws_dash["D27"].value is not None, "Dashboard C27 formula explanation missing"
    assert "BUY" in ws_dash["D31"].value, f"Expected decision rule formula in D31, got: {ws_dash['D31'].value}"
    
    # 3. Verify Assumptions values & formulas
    ws_ass = wb["Assumptions"]
    wacc_val = ws_ass["C20"].value
    assert isinstance(wacc_val, (int, float)) and wacc_val > 0, f"Expected numeric WACC in Assumptions, got: {wacc_val}"
    assert ws_ass["E20"].value is not None, "Expected formula explanation in Assumptions E20"
    
    # 4. Verify Pro-Forma Forecast populated numbers for all 5 projection years
    ws_pf = wb["Pro-Forma Forecast"]
    # Check Revenue row (row 6) has numbers in base and all 5 years (cols 3..8)
    for c_i in range(3, 9):
        val = ws_pf.cell(row=6, column=c_i).value
        assert isinstance(val, (int, float)) and val > 0, f"Expected numeric Revenue in Pro-Forma col {c_i}, got: {val}"
        
    # Check FCFF row (row 17) has numbers for all 5 forecast years (cols 4..8)
    for c_i in range(4, 9):
        fcff_val = ws_pf.cell(row=17, column=c_i).value
        assert isinstance(fcff_val, (int, float)), f"Expected numeric FCFF in Pro-Forma col {c_i}, got: {fcff_val}"
    assert ws_pf["I17"].value is not None, "Expected formula explanation for FCFF"
    
    # 5. Verify DCF Bridge pre-populated valuation waterfall
    ws_dcf = wb["DCF Bridge"]
    cum_pv = ws_dcf["C11"].value
    ev = ws_dcf["C24"].value
    eq_val = ws_dcf["C29"].value
    dcf_price = ws_dcf["C31"].value
    
    assert isinstance(cum_pv, (int, float)) and cum_pv > 0, f"Expected numeric cum PV of flows, got: {cum_pv}"
    assert isinstance(ev, (int, float)) and ev > 0, f"Expected numeric Enterprise Value, got: {ev}"
    assert isinstance(eq_val, (int, float)), f"Expected numeric Equity Value, got: {eq_val}"
    assert isinstance(dcf_price, (int, float)) and dcf_price > 0, f"Expected numeric DCF price, got: {dcf_price}"
    assert ws_dcf["D31"].value is not None, "Expected formula description for DCF price"
    
    # 6. Verify Relative Valuation populated medians and composite price
    ws_rel = wb["Relative Valuation"]
    # Check composite relative value row
    r_comp = ws_rel.max_row
    comp_val = ws_rel.cell(row=r_comp, column=7).value
    assert isinstance(comp_val, (int, float)) and comp_val > 0, f"Expected numeric composite relative value, got: {comp_val}"
    
    # 7. Verify Historical & Inputs opening balance sheet check
    ws_hist = wb["Historical & Inputs"]
    plug = ws_hist["C17"].value
    assert isinstance(plug, (int, float)) and abs(plug) <= 0.01, f"Expected 0.00 plug in Historical check, got: {plug}"
    assert "BALANCED" in str(ws_hist["C18"].value)
